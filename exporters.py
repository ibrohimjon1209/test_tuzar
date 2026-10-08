from __future__ import annotations

from pathlib import Path

from docx import Document


def export_questions_to_docx(questions: list[dict], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("Test savollari", level=1)

    for idx, item in enumerate(questions, start=1):
        doc.add_paragraph(f"{idx}. {item.get('question', '')}")
        for opt_index, option in enumerate(item.get('options', []), start=0):
            doc.add_paragraph(f"   {chr(65 + opt_index)}. {option}")
        doc.add_paragraph(f"To'g'ri javob: {chr(65 + int(item.get('correct_index', 0)))}")
        doc.add_paragraph(f"Tushuntirish: {item.get('explanation', '')}")
        doc.add_paragraph("-")

    doc.save(str(output_path))
    return output_path
